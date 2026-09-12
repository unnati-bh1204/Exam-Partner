import fitz
from docx import Document as DocxDocument
from pptx import Presentation

from app.extraction import extract_docx_sections, extract_pdf_pages, extract_pptx_slides


def test_extract_pdf_pages(tmp_path):
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Newton's First Law of Motion")
    doc.save(str(pdf_path))
    doc.close()

    pages = extract_pdf_pages(str(pdf_path))

    assert len(pages) == 1
    assert pages[0][0] == "page 1"
    assert "Newton" in pages[0][1]


def test_extract_pdf_skips_blank_pages(tmp_path):
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    doc.new_page()  # blank
    page = doc.new_page()
    page.insert_text((72, 72), "Chapter 2: Thermodynamics")
    doc.save(str(pdf_path))
    doc.close()

    pages = extract_pdf_pages(str(pdf_path))

    assert len(pages) == 1
    assert pages[0][0] == "page 2"


def test_extract_docx_sections(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("The mitochondria is the powerhouse of the cell.")
    doc.save(str(docx_path))

    sections = extract_docx_sections(str(docx_path))

    assert len(sections) == 1
    assert sections[0][0] == "document"
    assert "mitochondria" in sections[0][1]


def test_extract_pptx_slides(tmp_path):
    pptx_path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Photosynthesis"
    prs.save(str(pptx_path))

    slides = extract_pptx_slides(str(pptx_path))

    assert len(slides) == 1
    assert slides[0][0] == "slide 1"
    assert "Photosynthesis" in slides[0][1]
