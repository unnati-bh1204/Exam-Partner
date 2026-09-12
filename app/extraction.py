import base64

import fitz
from docx import Document as DocxDocument
from langchain_core.messages import HumanMessage
from langchain_google_genai import ChatGoogleGenerativeAI
from pptx import Presentation

from app.config import settings
from app.llm_response import text_from_response
from app.retry import with_retries


def extract_pdf_pages(path: str) -> list[tuple[str, str]]:
    doc = fitz.open(path)
    pages = []
    for i, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if text:
            pages.append((f"page {i}", text))
    doc.close()
    return pages


def extract_docx_sections(path: str) -> list[tuple[str, str]]:
    doc = DocxDocument(path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [("document", text)] if text else []


def extract_pptx_slides(path: str) -> list[tuple[str, str]]:
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        if texts:
            slides.append((f"slide {i}", "\n".join(texts)))
    return slides


def transcribe_image(path: str) -> list[tuple[str, str]]:
    with open(path, "rb") as f:
        image_data = base64.b64encode(f.read()).decode("utf-8")
    llm = ChatGoogleGenerativeAI(model="gemini-flash-latest", google_api_key=settings.gemini_api_key)
    message = HumanMessage(
        content=[
            {
                "type": "text",
                "text": (
                    "Transcribe all handwritten and printed text in this image exactly as written. "
                    "Return only the transcribed text, no commentary."
                ),
            },
            {"type": "image_url", "image_url": f"data:image/jpeg;base64,{image_data}"},
        ]
    )
    response = with_retries(lambda: llm.invoke([message]))
    text = text_from_response(response)
    return [("scanned page", text)] if text else []


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def extract_text(path: str, content_type: str) -> list[tuple[str, str]]:
    if content_type == "application/pdf":
        return extract_pdf_pages(path)
    if content_type == DOCX_MIME:
        return extract_docx_sections(path)
    if content_type == PPTX_MIME:
        return extract_pptx_slides(path)
    if content_type.startswith("image/"):
        return transcribe_image(path)
    raise ValueError(f"Unsupported file type: {content_type}")
